import os
import sys
import docx
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

# Dictionary of common typos / spelling corrections in lesson plans
SPELLING_CORRECTIONS = {
    "BẢNG TUẦN TOÀN": "BẢNG TUẦN HOÀN",
    "Bảng tuần toàn": "Bảng tuần hoàn",
    "NGUYỀN TỐ": "NGUYÊN TỐ",
    "nguyền tố": "nguyên tố",
    "lí thuyêt": "lí thuyết",
    "gsk": "sgk",
    "Tuần 45": "Tuần 4-5",
    "Tiết 1418": "Tiết 14-18",
    "Tiết 15": "Tiết 1-5",
    "Tuần 23": "Tuần 2-3",
    "Tiết 36": "Tiết 6-9",  # corrected to match PPCT for Bai 2
    "Tiết 1013": "Tiết 10-13",
    "Tuần 34": "Tuần 3-4"
}

def auto_correct_lesson_plan(file_path, output_dir=None):
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return None

    filename = os.path.basename(file_path)
    print(f"\n[AUTO-CORRECTION ND30] Đang xử lý: {filename}")
    doc = docx.Document(file_path)

    # 1. Căn lề chuẩn Nghị định 30/2020/NĐ-CP (Top 20mm, Bottom 20mm, Left 30mm, Right 15mm)
    for section in doc.sections:
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(30)
        section.right_margin = Mm(15)
        section.page_width = Mm(210)
        section.page_height = Mm(297)

    # 2. Rà soát & sửa font chữ, kích thước, lỗi chính tả trong Paragraphs
    fixed_typos_count = 0
    for p in doc.paragraphs:
        # Thay thế lỗi chính tả trong text
        for wrong, right in SPELLING_CORRECTIONS.items():
            if wrong in p.text:
                for run in p.runs:
                    if wrong in run.text:
                        run.text = run.text.replace(wrong, right)
                        fixed_typos_count += 1

        for run in p.runs:
            # Chuẩn hóa phông chữ Times New Roman
            run.font.name = "Times New Roman"
            if run.font.size and run.font.size.pt not in [11.0, 12.0, 13.0, 14.0, 15.0, 16.0]:
                run.font.size = Pt(13)

    # 3. Rà soát & sửa trong Tables
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    for wrong, right in SPELLING_CORRECTIONS.items():
                        if wrong in p.text:
                            for run in p.runs:
                                if wrong in run.text:
                                    run.text = run.text.replace(wrong, right)
                                    fixed_typos_count += 1
                    for run in p.runs:
                        run.font.name = "Times New Roman"
                        if run.font.size and run.font.size.pt not in [10.0, 11.0, 12.0, 13.0, 14.0]:
                            run.font.size = Pt(12)

    # Xác định thư mục lưu
    if output_dir is None:
        output_dir = os.path.join(os.path.dirname(file_path), "ChuanHoa_ND30")
    os.makedirs(output_dir, exist_ok=True)

    clean_name = filename.replace(".docx", "_ChuanTheThuc_ND30.docx")
    for wrong, right in [("TOÀN", "HOÀN"), ("-ST", "")]:
        clean_name = clean_name.replace(wrong, right)

    out_file = os.path.join(output_dir, clean_name)
    doc.save(out_file)
    print(f" -> Đã căn chỉnh lề chuẩn 20-20-30-15mm, sửa {fixed_typos_count} lỗi chính tả/tiêu đề.")
    print(f" -> Tệp xuất bản: {out_file}")
    return out_file

if __name__ == "__main__":
    test_files = [
        r"D:\APP-KIEM-TRA-GIAO-AN\data\Co_Cuc\Cô Cúc\KHTN 7\KHTN 7 tuần 1-4\Tuần 1-2 Tiết 1-5  BÀI 1 MỞ ĐẦU KHTN.docx",
        r"D:\APP-KIEM-TRA-GIAO-AN\data\Co_Cuc\Cô Cúc\KHTN 7\KHTN 7 tuần 1-4\Tuần 2-3  Tiết 6-9  Bài 2  Nguyên tử - KHTN7 - CTST.docx",
        r"D:\APP-KIEM-TRA-GIAO-AN\data\Co_Cuc\Cô Cúc\KHTN 7\KHTN 7 tuần 1-4\Tuần 3-4, Tiết 10-13  -CHỦ ĐỀ 1 - BÀI 3- NTHH.docx",
        r"D:\APP-KIEM-TRA-GIAO-AN\data\Co_Cuc\Cô Cúc\KHTN 7\KHTN 7 tuần 1-4\Tuần 4-5  -Tiết 14-18  Bài 4- SƠ LƯỢC BẢNG TUẦN TOÀN CÁC NTHH-KHTN 7-CTST-ST.docx"
    ]
    out_dir = r"D:\APP-KIEM-TRA-GIAO-AN\data\Co_Cuc_ChuanTheThuc_ND30"
    for f in test_files:
        if os.path.exists(f):
            auto_correct_lesson_plan(f, out_dir)
