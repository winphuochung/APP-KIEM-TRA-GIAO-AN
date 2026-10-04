import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_table_borders(table, color="7F7F7F", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def generate_inspection_report(teacher_name="Phạm Thị Cúc", inspector_name="Lê Văn Thắng", period_str="Tuần 1 đến tuần 4", date_str="ngày 03 tháng 10 năm 2026", output_path=None):
    if output_path is None:
        output_path = r"D:\APP-KIEM-TRA-GIAO-AN\BIEN_BAN_KIEM_TRA_GIAO_AN.docx"

    doc = docx.Document()
    
    # Nghị định 30/2020/NĐ-CP: Top=20mm, Bottom=20mm, Left=30mm, Right=15mm
    for section in doc.sections:
        section.top_margin = docx.shared.Mm(20)
        section.bottom_margin = docx.shared.Mm(20)
        section.left_margin = docx.shared.Mm(30)
        section.right_margin = docx.shared.Mm(15)
        section.page_width = docx.shared.Mm(210)
        section.page_height = docx.shared.Mm(297)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(4)

    # Header Table
    t0 = doc.add_table(rows=3, cols=2)
    t0.alignment = WD_TABLE_ALIGNMENT.CENTER
    t0.autofit = False
    
    col_widths_0 = [docx.shared.Mm(75), docx.shared.Mm(90)]
    for row in t0.rows:
        for idx, width in enumerate(col_widths_0):
            row.cells[idx].width = width

    # Row 0
    c00 = t0.rows[0].cells[0].paragraphs[0]
    c00.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c00.add_run("TRƯỜNG TH & THCS PHƯỚC HƯNG")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    c01 = t0.rows[0].cells[1].paragraphs[0]
    c01.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c01.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True

    # Row 1
    c10 = t0.rows[1].cells[0].paragraphs[0]
    c10.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c10.add_run("TỔ: TOÁN – KHTN - CN")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True

    c11 = t0.rows[1].cells[1].paragraphs[0]
    c11.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c11.add_run("Độc lập - Tự do - Hạnh phúc")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.underline = True

    # Row 2
    c20 = t0.rows[2].cells[0].paragraphs[0]
    c20.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c20.add_run("Số: 02/BB-TCM")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.italic = True

    c21 = t0.rows[2].cells[1].paragraphs[0]
    c21.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c21.add_run(f"Nhơn Hội, {date_str}")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.italic = True

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(6)
    r_t1 = p_title.add_run("BIÊN BẢN KIỂM TRA ĐỊNH KỲ HỒ SƠ KẾ HOẠCH BÀI DẠY\n")
    r_t1.font.bold = True
    r_t1.font.size = Pt(14)
    r_t1.font.color.rgb = RGBColor(0, 32, 96)
    r_t2 = p_title.add_run(f"(Giai đoạn: {period_str} – Năm học: 2026 – 2027)")
    r_t2.font.bold = True
    r_t2.font.italic = True
    r_t2.font.size = Pt(13)

    # 1. Thông tin chung
    p1 = doc.add_paragraph()
    r = p1.add_run("1. Thông tin chung:")
    r.font.bold = True
    r.font.size = Pt(13)

    info_items = [
        ("Người kiểm tra:", f" {inspector_name} – Tổ trưởng chuyên môn Tổ Toán – KHTN – CN."),
        ("Người được kiểm tra:", f" {teacher_name} – Giáo viên thuộc Tổ chuyên môn Toán – KHTN – CN."),
        ("Giai đoạn kiểm tra:", f" {period_str}."),
        ("Thời gian kiểm tra:", f" 08 giờ 30 phút, {date_str}."),
        ("Địa điểm kiểm tra:", " Văn phòng Tổ chuyên môn Toán – KHTN – CN, Trường TH & THCS Phước Hưng."),
        ("Nội dung kiểm tra:", f" Kiểm tra kế hoạch bài dạy (giáo án) định kỳ {period_str} đối chiếu Công văn 5512/BGDĐT, Nghị định 30/2020/NĐ-CP và Phân phối chương trình nhà trường.")
    ]
    for label, val in info_items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(2)
        r_lbl = p.add_run(f"• {label}")
        r_lbl.font.bold = True
        p.add_run(val)

    # 2. Bảng tổng hợp
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(6)
    r = p2.add_run("2. Bảng tổng hợp và nhận xét chi tiết các kế hoạch bài dạy được kiểm tra:")
    r.font.bold = True
    r.font.size = Pt(13)

    t1 = doc.add_table(rows=1, cols=7)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t1.autofit = False
    set_table_borders(t1, color="7F7F7F", sz="4")

    headers = ['STT', 'Môn/Lớp', 'Tuần', 'Tiết (PPCT)', 'Tên bài dạy / Chủ đề', 'Nhận xét chi tiết (5512, Bloom, Thể thức, Smart Comments)', 'Xếp loại']
    widths = [docx.shared.Mm(10), docx.shared.Mm(18), docx.shared.Mm(14), docx.shared.Mm(16), docx.shared.Mm(32), docx.shared.Mm(65), docx.shared.Mm(15)]

    for i, title in enumerate(headers):
        cell = t1.rows[0].cells[i]
        cell.width = widths[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_background(cell, "1F4E79")
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        r = p.add_run(title)
        r.font.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(255, 255, 255)

    data = [
        (
            "1", "KHTN 7\n(CTST)", "Tuần 1, 2", "Tiết 1, 2, 3, 4, 5",
            "Bài 1: Mở đầu - Phương pháp và kĩ năng học tập môn KHTN\n(5 tiết)",
            [
                ("Tiến trình 5512: ", "Đầy đủ chuỗi 4 hoạt động. Phân định rõ 4 bước trong tổ chức dạy học."),
                ("Nghiệp vụ sư phạm & Bloom: ", "Nội dung bám sát SGK CTST; hình thành tốt các kĩ năng tiến trình khoa học (quan sát, phân loại, đo đạc, dự đoán, viết báo cáo). Mức độ nhận thức phân bố hợp lý."),
                ("Kỹ thuật & Thể thức (NĐ 30): ", "Căn lề chuẩn (Trên 20mm, Dưới 20mm, Trái 30mm, Phải 15mm), font Times New Roman 13-14pt. Lỗi nhỏ: Tiêu đề thiếu gạch nối 'Tiết 15' (sửa thành 'Tiết 1-5')."),
                ("Smart Comments: ", "[Ghi nhận điểm sáng] Hệ thống câu hỏi dẫn dắt HS làm quen với nghiên cứu thực nghiệm rất sinh động. [Góp ý cải tiến] Hoạt động Vận dụng nên chuyển thành bài tập dự án thực hành nhỏ tại nhà (như đo đạc kích thước thực tế) để học sinh quay video ngắn nộp lại.")
            ],
            "Tốt"
        ),
        (
            "2", "KHTN 7\n(CTST)", "Tuần 2, 3", "Tiết 6, 7, 8, 9",
            "Chủ đề 1:\nBài 2: Nguyên tử\n(4 tiết)",
            [
                ("Tiến trình 5512: ", "Cấu trúc rõ ràng, logic. Mục tiêu bám sát năng lực nhận thức KHTN về mô hình nguyên tử Rutherford - Bohr."),
                ("Nghiệp vụ sư phạm & Bloom: ", "Khoa học, chuẩn xác. Dẫn dắt học sinh khám phá cấu tạo hạt nhân (p, n) và vỏ electron đạt tốt mức Thông hiểu và Vận dụng."),
                ("Kỹ thuật & Thể thức (NĐ 30): ", "Căn lề đạt chuẩn 20-20-30-15mm. [TỒN TẠI KỸ THUẬT]: Tiêu đề ghi 'Tiết 36' lệch so với PPCT nhà trường (PPCT là Tiết 6→9)."),
                ("Smart Comments: ", "[Ghi nhận điểm sáng] Sử dụng hình ảnh trực quan mô phỏng chuyển động electron rất tốt. [Góp ý cải tiến] Ở Hoạt động 2, Thầy/Cô cần bổ sung khung 'Chốt kiến thức cốt lõi' riêng biệt để học sinh dễ ghi nhớ bài học.")
            ],
            "Tốt"
        ),
        (
            "3", "KHTN 7\n(CTST)", "Tuần 3, 4", "Tiết 10, 11, 12, 13",
            "Chủ đề 1:\nBài 3: Nguyên tố hóa học\n(4 tiết)",
            [
                ("Tiến trình 5512: ", "Có đầy đủ các hoạt động học tập. Chuỗi bài tập nhận biết ký hiệu hóa học và tính nguyên tử khối phong phú."),
                ("Nghiệp vụ sư phạm & Bloom: ", "Nội dung chuẩn xác. Rèn luyện tốt kĩ năng tra cứu bảng tuần hoàn ở mức Nhận biết và Vận dụng thấp."),
                ("Kỹ thuật & Thể thức (NĐ 30): ", "[KHÔNG ĐẠT THỂ THỨC]: Tệp định dạng .doc cũ; căn lề trái chỉ 25mm (chuẩn 30-35mm); tiêu đề ghi 'NGUYỀN TỐ' (sai dấu); ghi nhầm '3 tiết' thay vì 4 tiết."),
                ("Smart Comments: ", "[Ghi nhận điểm sáng] Hệ thống phiếu bài tập rèn kỹ năng viết kí hiệu hóa học rất tỉ mỉ. [Góp ý cải tiến] Phần Khởi động nên dùng trò chơi tương tác 'Đố vui hóa học'; khẩn trương chuyển sang .docx và căn lề trái 30mm theo đúng NĐ 30.")
            ],
            "Khá"
        ),
        (
            "4", "KHTN 7\n(CTST)", "Tuần 4, 5", "Tiết 14, 15, 16 (tuần 4)\n& 17, 18",
            "Chủ đề 1:\nBài 4: Sơ lược về bảng tuần hoàn các nguyên tố hóa học\n(5 tiết)",
            [
                ("Tiến trình 5512: ", "Được thiết kế công phu, tích hợp video tư liệu lịch sử bảng tuần hoàn Mendeleev."),
                ("Nghiệp vụ sư phạm & Bloom: ", "Khai thác sâu sắc cấu trúc ô nguyên tố, chu kì và nhóm; phát triển tốt tư duy so sánh và suy luận quy luật tuần hoàn."),
                ("Kỹ thuật & Thể thức (NĐ 30): ", "[LỖI THỂ THỨC & CHÍNH TẢ]: Căn lề trái chỉ 20mm (vi phạm chuẩn 30-35mm); tiêu đề sai chính tả 'BẢNG TUẦN TOÀN'."),
                ("Smart Comments: ", "[Ghi nhận điểm sáng] Tích hợp CNTT xuất sắc với video tư liệu kích thích hứng thú. [Góp ý cải tiến] Tại Hoạt động luyện tập, nên tổ chức mini-game 'Truy tìm nguyên tố bí mật' theo tọa độ chu kỳ - nhóm để phân hóa học sinh khá giỏi.")
            ],
            "Tốt"
        ),
        (
            "5", "KHTN 9\n(KNTT)", "Tuần 2", "Tiết 1, 2\n(Sinh học)",
            "Bài 36: Khái quát về di truyền học\n(2 tiết)",
            [
                ("Tiến trình 5512: ", "Xuất sắc. Bổ sung hệ thống Rubrics đánh giá theo tiêu chí và mức độ rất bài bản."),
                ("Nghiệp vụ sư phạm & Bloom: ", "Làm rõ mối quan hệ di truyền - biến dị; khẳng định vai trò trung tâm của gene và công lao của Mendel. Đạt chuẩn mức Vận dụng giải thích hiện tượng đời sống."),
                ("Kỹ thuật & Thể thức (NĐ 30): ", "Đạt chuẩn xuất sắc: Căn lề 20-20-30-20mm đúng NĐ 30, font Times New Roman thống nhất, không lỗi chính tả."),
                ("Smart Comments: ", "[Ghi nhận điểm sáng] Giáo án mẫu mực, có Rubric đánh giá năng lực đặc thù sinh học rất đáng nhân rộng. [Góp ý cải tiến] Ở Hoạt động Vận dụng, có thể gợi mở học sinh vẽ sơ đồ phả hệ tính trạng đơn giản trong gia đình mình.")
            ],
            "Tốt"
        ),
        (
            "6", "KHTN 9\n(KNTT)", "Tuần 3, 4, 5, 6", "Tiết 3, 4 (tuần 3, 4)\n& Tiết 5, 6",
            "Bài 37: Các quy luật di truyền của Menđen\n(4 tiết)",
            [
                ("Tiến trình 5512: ", "Đầy đủ các hoạt động theo phân phối 4 tiết. Đã soạn đón đầu tiến độ các tuần kế tiếp."),
                ("Nghiệp vụ sư phạm & Bloom: ", "Chuyên môn vững vàng. Phân tích kết quả thí nghiệm lai một và hai tính trạng chặt chẽ. Xây dựng sơ đồ lai chuẩn mực, đạt mức độ Vận dụng cao."),
                ("Kỹ thuật & Thể thức (NĐ 30): ", "Căn lề chuẩn (20-20-30-15mm), bảng biểu khoa học, phông chữ đồng nhất."),
                ("Smart Comments: ", "[Ghi nhận điểm sáng] Phiếu học tập số 1 và số 2 thiết kế khoa học, định hướng HS tự phát hiện quy luật. [Góp ý cải tiến] Ở nội dung Lai phân tích, cần mở rộng nhận định về ứng dụng trong kiểm tra độ thuần chủng giống cây trồng/vật nuôi.")
            ],
            "Tốt"
        )
    ]

    for item in data:
        row_cells = t1.add_row().cells
        for col_idx in range(7):
            row_cells[col_idx].width = widths[col_idx]
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=80, right=80)
            row_cells[col_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row_cells[0].paragraphs[0].add_run(item[0]).font.size = Pt(11)

        row_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = row_cells[1].paragraphs[0].add_run(item[1])
        r.font.bold = True
        r.font.size = Pt(11)

        row_cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row_cells[2].paragraphs[0].add_run(item[2]).font.size = Pt(11)

        row_cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row_cells[3].paragraphs[0].add_run(item[3]).font.size = Pt(11)

        r = row_cells[4].paragraphs[0].add_run(item[4])
        r.font.bold = True
        r.font.size = Pt(11)

        p_cmt = row_cells[5].paragraphs[0]
        for idx_sub, (sub_lbl, sub_txt) in enumerate(item[5]):
            p_sub = p_cmt if idx_sub == 0 else row_cells[5].add_paragraph()
            p_sub.paragraph_format.space_after = Pt(2)
            p_sub.paragraph_format.line_spacing = 1.1
            r_sub_lbl = p_sub.add_run(f"• {sub_lbl}")
            r_sub_lbl.font.bold = True
            r_sub_lbl.font.size = Pt(10.5)
            if "KHÔNG ĐẠT" in sub_lbl or "TỒN TẠI" in sub_txt or "LỖI" in sub_lbl:
                r_sub_lbl.font.color.rgb = RGBColor(192, 0, 0)
            elif "Smart Comments" in sub_lbl:
                r_sub_lbl.font.color.rgb = RGBColor(0, 112, 192)
            p_sub.add_run(sub_txt).font.size = Pt(10.5)

        row_cells[6].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = row_cells[6].paragraphs[0].add_run(item[6])
        r.font.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0, 128, 0) if item[6] == "Tốt" else RGBColor(192, 112, 0)

    # 3. Đánh giá chung
    p3 = doc.add_paragraph()
    p3.paragraph_format.space_before = Pt(8)
    p3.add_run("3. Đánh giá chung:").font.bold = True

    doc.add_paragraph("a) Ưu điểm:").runs[0].font.bold = True
    for u in [
        "Tiến độ: Hoàn thành 100% kế hoạch bài dạy từ Tuần 1 đến 4, có sự chuẩn bị đón đầu sang tuần 5 và 6.",
        "Sư phạm (5512): Tuân thủ đầy đủ chuỗi 4 hoạt động và cấu trúc mục tiêu theo định hướng GDPT 2018.",
        "Đổi mới kiểm tra đánh giá: Điểm sáng mẫu mực là hệ thống Rubrics 3 mức độ ở môn KHTN 9."
    ]:
        p = doc.add_paragraph(f"• {u}")
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(2)

    doc.add_paragraph("b) Tồn tại cần khắc phục:").runs[0].font.bold = True
    for t in [
        "Thể thức: Cần chỉnh lại lề trái Bài 3 và 4 thành 30mm theo đúng Nghị định 30/2020/NĐ-CP.",
        "Chính tả & Thống kê: Khắc phục lỗi chính tả ở tiêu đề (TOÀN -> HOÀN, NGUYỀN -> NGUYÊN) và chỉnh khớp số tiết PPCT."
    ]:
        p = doc.add_paragraph(f"• {t}")
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(2)

    # 4. Đề xuất
    p4 = doc.add_paragraph()
    p4.paragraph_format.space_before = Pt(6)
    p4.add_run("4. Đề xuất, kiến nghị:").font.bold = True
    for kn in [
        "Sử dụng công cụ Auto-Correction để chuẩn hóa thể thức toàn bộ tệp giáo án trước khi lưu trữ đám mây.",
        "Nhân rộng phương pháp xây dựng tiêu chí Rubrics sang các phân môn Hóa học và Vật lý."
    ]:
        p = doc.add_paragraph(f"• {kn}")
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(2)

    # 5. Kết luận
    p5 = doc.add_paragraph()
    p5.paragraph_format.space_before = Pt(6)
    p5.add_run("5. Kết luận của người kiểm tra:").font.bold = True
    p_kl = doc.add_paragraph(f"• Hồ sơ kế hoạch bài dạy của cô {teacher_name} đạt loại: ")
    p_kl.paragraph_format.left_indent = Inches(0.2)
    r_xl = p_kl.add_run("TỐT")
    r_xl.font.bold = True
    r_xl.font.color.rgb = RGBColor(0, 128, 0)
    p_kl.add_run(". Ghi nhận sự chuẩn bị chu đáo, chuyên nghiệp và tinh thần đổi mới sáng tạo.")

    # Signatures
    doc.add_paragraph().paragraph_format.space_before = Pt(10)
    t2 = doc.add_table(rows=2, cols=2)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2.autofit = False
    for row in t2.rows:
        row.cells[0].width = docx.shared.Mm(80)
        row.cells[1].width = docx.shared.Mm(85)

    c0 = t2.rows[0].cells[0].paragraphs[0]
    c0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c0.add_run("NGƯỜI ĐƯỢC KIỂM TRA").font.bold = True

    c1 = t2.rows[0].cells[1].paragraphs[0]
    c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c1.add_run("NGƯỜI KIỂM TRA\nTỔ TRƯỞNG CHUYÊN MÔN").font.bold = True

    p_sig0 = t2.rows[1].cells[0].paragraphs[0]
    p_sig0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sig0.paragraph_format.space_before = Pt(45)
    p_sig0.add_run(teacher_name).font.bold = True

    p_sig1 = t2.rows[1].cells[1].paragraphs[0]
    p_sig1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sig1.paragraph_format.space_before = Pt(45)
    p_sig1.add_run(inspector_name).font.bold = True

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path
